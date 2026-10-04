"""
HisseRadarPro — Excel export
============================
- build_model_workbook: the DCF / justified P/B model as a live workbook. Every projection cell is a
  formula that points at the "Varsayımlar" sheet, so changing one assumption in Excel rebuilds the
  model (the sensitivity grid is a static snapshot and says so).
- build_table_workbook: generic sheets of rows for screener, backtest, strategy and portfolio exports.
"""

import datetime
import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GOLD = "C8A24A"
HEAD_FILL = PatternFill("solid", fgColor="121214")
INPUT_FILL = PatternFill("solid", fgColor="FFF4D6")
HEAD_FONT = Font(bold=True, color=GOLD)
TITLE_FONT = Font(bold=True, size=14)
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
PCT = "0.0%"
NUM = "#,##0.00"
BIG = "#,##0"


def _title(ws, text: str, sub: str = ""):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="666666")


def _header(ws, row: int, labels: List[str]):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: List[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_model_workbook(m: Dict[str, Any], use_thesis: bool = False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Varsayımlar"
    ticker = m["ticker"]
    model = m.get("thesis_model") if use_thesis and m.get("thesis_model") else m["base_model"]
    inputs = (m.get("thesis_inputs") if use_thesis and m.get("thesis_inputs") else m["inputs"]) or {}
    bal = m["balance"]
    src = m.get("sources", {})
    _title(ws, f"{ticker} · {'DCF modeli' if m['model'] == 'dcf' else 'Hak edilen PD/DD modeli'}",
           f"HisseRadar Pro · {datetime.date.today().isoformat()} · Sarı hücreler girdidir; değiştirince model yeniden hesaplanır.")
    _header(ws, 4, ["Girdi", "Değer", "Kaynak"])
    _widths(ws, [34, 18, 90])

    # (key, label, value, format, source key)
    cc = model["cost_of_capital"]
    rows = [
        ("price", "Fiyat (TL)", m["price"], NUM, None),
        ("shares", "Hisse sayısı", bal["shares"], BIG, "shares"),
        ("risk_free", "Risksiz getiri", inputs.get("risk_free"), PCT, "risk_free"),
        ("erp", "Hisse senedi risk primi", inputs.get("equity_premium"), PCT, "equity_premium"),
        ("beta", "Beta", inputs.get("beta"), "0.00", "beta"),
        ("g_term", "Uç büyüme", inputs.get("terminal_growth"), PCT, "terminal_growth"),
    ]
    if m["model"] == "dcf":
        rows += [
            ("revenue", "Son 12 ay gelir (TL)", bal["revenue_ttm"], BIG, "revenue_ttm"),
            ("g1", "İlk yıl gelir büyümesi", inputs.get("growth_y1"), PCT, "growth_y1"),
            ("margin", "Faaliyet (FVÖK) marjı", inputs.get("ebit_margin"), PCT, "ebit_margin"),
            ("tax", "Vergi oranı", inputs.get("tax_rate"), PCT, "tax_rate"),
            ("da", "Amortisman / gelir", inputs.get("da_pct"), PCT, "da_pct"),
            ("capex0", "Yatırım harcaması / gelir (başlangıç)", inputs.get("capex_pct"), PCT, "capex_pct"),
            ("capexT", "Yatırım harcaması / gelir (10. yıl)", (inputs.get("da_pct") or 0) + 0.01, PCT, "capex_terminal_pct"),
            ("nwc", "İşletme sermayesi / gelir artışı", 0.10, PCT, "nwc_pct_of_delta_rev"),
            ("spread", "Borç maliyeti farkı", 0.025, PCT, None),
            ("debt", "Toplam borç (TL)", bal["total_debt"], BIG, "total_debt"),
            ("cash", "Nakit (TL)", bal["cash"], BIG, "cash"),
            ("minority", "Azınlık payları (TL)", bal["minority"], BIG, None),
            ("mcap", "Piyasa değeri (TL)", bal["market_cap"], BIG, None),
        ]
    else:
        rows += [
            ("roe", "Sürdürülebilir özkaynak kârlılığı", inputs.get("roe"), PCT, "roe"),
            ("book", "Defter değeri (TL)", bal["book_equity"], BIG, "book_equity"),
        ]
    ref: Dict[str, str] = {}
    r = 5
    for key, label, value, fmt, sk in rows:
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=value)
        c.number_format = fmt
        c.fill = INPUT_FILL
        ws.cell(row=r, column=3, value=src.get(sk, "") if sk else "")
        ref[key] = f"Varsayımlar!$B${r}"
        r += 1

    # Derived cost of capital (formulas)
    r += 1
    ws.cell(row=r, column=1, value="Özkaynak maliyeti").font = Font(bold=True)
    ws.cell(row=r, column=2, value=f"={ref['risk_free']}+{ref['beta']}*{ref['erp']}").number_format = PCT
    ref["coe"] = f"Varsayımlar!$B${r}"
    r += 1
    if m["model"] == "dcf":
        ws.cell(row=r, column=1, value="Borç ağırlığı")
        ws.cell(row=r, column=2, value=f"={ref['debt']}/({ref['debt']}+{ref['mcap']})").number_format = PCT
        ref["wd"] = f"Varsayımlar!$B${r}"
        r += 1
        ws.cell(row=r, column=1, value="AOSM (WACC)").font = Font(bold=True)
        ws.cell(row=r, column=2, value=f"={ref['coe']}*(1-{ref['wd']})+({ref['risk_free']}+{ref['spread']})*(1-{ref['tax']})*{ref['wd']}").number_format = PCT
        ref["wacc"] = f"Varsayımlar!$B${r}"

        # Projection sheet
        d = wb.create_sheet("DCF")
        _title(d, f"{ticker} · 10 yıllık serbest nakit akımı projeksiyonu", "Tüm hücreler formüldür; değerler bugünkü TL'dir (reel).")
        labels = ["Kalem"] + [f"Yıl {t}" for t in range(1, 11)]
        _header(d, 4, labels)
        _widths(d, [30] + [16] * 10)
        lines = ["Büyüme", "Gelir", "FVÖK", "NOPAT", "Amortisman", "Yatırım / gelir", "Yatırım harcaması", "İşletme sermayesi artışı",
                 "Serbest nakit akımı (FCFF)", "İskonto faktörü", "Bugünkü değer"]
        for i, name in enumerate(lines):
            d.cell(row=5 + i, column=1, value=name).font = Font(bold=name in ("Serbest nakit akımı (FCFF)", "Bugünkü değer"))
        for t in range(1, 11):
            col = get_column_letter(t + 1)
            prev_rev = ref["revenue"] if t == 1 else f"{get_column_letter(t)}6"
            d[f"{col}5"] = f"={ref['g1']}+({ref['g_term']}-{ref['g1']})*({t}-1)/9"
            d[f"{col}6"] = f"={prev_rev}*(1+{col}5)"
            d[f"{col}7"] = f"={col}6*{ref['margin']}"
            d[f"{col}8"] = f"={col}7*(1-{ref['tax']})"
            d[f"{col}9"] = f"={col}6*{ref['da']}"
            d[f"{col}10"] = f"={ref['capex0']}+({ref['capexT']}-{ref['capex0']})*({t}-1)/9"
            d[f"{col}11"] = f"={col}6*{col}10"
            d[f"{col}12"] = f"=({col}6-{prev_rev})*{ref['nwc']}"
            d[f"{col}13"] = f"={col}8+{col}9-{col}11-{col}12"
            d[f"{col}14"] = f"=1/(1+{ref['wacc']})^{t}"
            d[f"{col}15"] = f"={col}13*{col}14"
            for rr, fmt in ((5, PCT), (10, PCT), (14, "0.0000")):
                d[f"{col}{rr}"].number_format = fmt
            for rr in (6, 7, 8, 9, 11, 12, 13, 15):
                d[f"{col}{rr}"].number_format = BIG
        s = 18
        out = [
            ("Projeksiyonun bugünkü değeri", "=SUM(B15:K15)", BIG),
            ("Uç değer", f"=K13*(1+{ref['g_term']})/({ref['wacc']}-{ref['g_term']})", BIG),
            ("Uç değerin bugünkü değeri", f"=B{s + 1}*K14", BIG),
            ("Firma değeri", f"=B{s}+B{s + 2}", BIG),
            ("Özkaynak değeri", f"=B{s + 3}-{ref['debt']}+{ref['cash']}-{ref['minority']}", BIG),
            ("Hisse başı değer (TL)", f"=B{s + 4}/{ref['shares']}", NUM),
            ("Fiyata göre fark", f"=B{s + 5}/{ref['price']}-1", PCT),
        ]
        for i, (label, formula, fmt) in enumerate(out):
            d.cell(row=s + i, column=1, value=label).font = Font(bold=label.startswith("Hisse"))
            c = d.cell(row=s + i, column=2, value=formula)
            c.number_format = fmt
    else:
        pbs = wb.create_sheet("PD-DD modeli")
        _title(pbs, f"{ticker} · Hak edilen PD/DD", "PD/DD = (ROE − g) / (Özkaynak maliyeti − g)")
        _widths(pbs, [36, 18])
        out = [
            ("Hak edilen PD/DD", f"=({ref['roe']}-{ref['g_term']})/({ref['coe']}-{ref['g_term']})", "0.00"),
            ("Hisse başı defter değeri (TL)", f"={ref['book']}/{ref['shares']}", NUM),
            ("Hisse başı değer (TL)", "=MAX(0,B4*B5)", NUM),
            ("Fiyata göre fark", f"=B6/{ref['price']}-1", PCT),
        ]
        for i, (label, formula, fmt) in enumerate(out):
            pbs.cell(row=4 + i, column=1, value=label)
            c = pbs.cell(row=4 + i, column=2, value=formula)
            c.number_format = fmt

    # Sensitivity snapshot
    sens = model["sensitivity"]
    sw = wb.create_sheet("Duyarlılık")
    _title(sw, f"{sens['rows_label']} × {sens['cols_label']} → hisse başı değer", "Statik görüntü: Excel'de girdileri değiştirince bu tablo güncellenmez.")
    sw.cell(row=4, column=1, value=f"{sens['rows_label']} \\ {sens['cols_label']}").font = HEAD_FONT
    sw.cell(row=4, column=1).fill = HEAD_FILL
    for j, g in enumerate(sens["cols"]):
        c = sw.cell(row=4, column=2 + j, value=g)
        c.number_format = PCT
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for i, w in enumerate(sens["rows"]):
        c = sw.cell(row=5 + i, column=1, value=w)
        c.number_format = PCT
        c.font = Font(bold=True)
        for j, v in enumerate(sens["values"][i]):
            sw.cell(row=5 + i, column=2 + j, value=v).number_format = NUM
    _widths(sw, [24] + [14] * len(sens["cols"]))

    notes = wb.create_sheet("Notlar")
    _widths(notes, [120])
    lines = [
        "Bu model basitleştirilmiştir ve yatırım tavsiyesi değildir.",
        "TL tutarları piyasa verisinden türetilir (piyasa değeri, F/K, PD/DD); finansal tablolardan yalnızca oranlar kullanılır.",
        "BIST şirketleri TMS 29 enflasyon muhasebesiyle raporladığı için DCF reel (bugünkü TL) bazdadır; bankalar TMS 29 uygulamadığından PD/DD modeli nominaldir.",
        "TMS 29 faaliyet marjları olağan dışı düşük görünebilir; ters model (fiyatın ima ettiği marj/büyüme) bu yüzden önemlidir.",
    ]
    if model.get("reverse"):
        rv = model["reverse"]
        if rv.get("value") is not None:
            lines.append(f"{rv['label']}: {rv['value']:.1%} (modeldeki: {rv['current']:.1%})" if rv.get("current") is not None else f"{rv['label']}: {rv['value']:.1%}")
        if rv.get("margin_value") is not None:
            lines.append(f"{rv['margin_label']}: {rv['margin_value']:.1%} (modeldeki: {rv['margin_current']:.1%})")
    for i, line in enumerate(lines, start=1):
        notes.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True)
    return _bytes(wb)


def build_table_workbook(sheets: List[Dict[str, Any]], title: str = "") -> bytes:
    """sheets: [{name, columns: [{key,label,format?}], rows: [dict]}]"""
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets[:10]:
        ws = wb.create_sheet((sh.get("name") or "Sayfa")[:31])
        cols = sh.get("columns") or []
        start = 1
        if title:
            _title(ws, title, f"HisseRadar Pro · {datetime.date.today().isoformat()} · Yatırım tavsiyesi değildir.")
            start = 4
        _header(ws, start, [c.get("label") or c["key"] for c in cols])
        fmt_map = {"pct": PCT, "num": NUM, "int": BIG, "pct100": '0.0"%"'}
        for i, row in enumerate(sh.get("rows") or [], start=start + 1):
            for j, c in enumerate(cols, start=1):
                v = row.get(c["key"]) if isinstance(row, dict) else None
                if isinstance(v, float) and v != v:  # NaN
                    v = None
                cell = ws.cell(row=i, column=j, value=v)
                if c.get("format") in fmt_map:
                    cell.number_format = fmt_map[c["format"]]
        _widths(ws, [max(10, min(40, len(str(c.get("label") or c["key"])) + 6)) for c in cols])
        ws.freeze_panes = ws.cell(row=start + 1, column=2)
    return _bytes(wb)
